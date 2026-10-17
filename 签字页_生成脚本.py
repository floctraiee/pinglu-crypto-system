# -*- coding: utf-8 -*-
"""按官方《审批及签字页》模板生成签字页。

保持模板版式与说明文字不变，只填入作品题目与三人任务分工；
姓名、院系/部门、专业、签字与作品编号留空，打印后手写补齐（匿名与实名信息按官方要求处理）。
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(r"C:\Users\67\AppData\Local\Temp\附件三---审批及签字页.docx")
OUT = ROOT / "审批及签字页_平陆运河国密可信采集与篡改溯源系统.docx"

TITLE = "平陆运河多源感知数据国密可信采集与篡改溯源系统"

# 按在作品中的贡献大小排序：队长 → 队员A → 队员B
TASKS = [
    "中心平台、三端联调、性能实验与报告提交",
    "国密密码模块、七类设备模拟与批次签名",
    "边缘网关、离线补传与异常样本验证",
]


def set_run(run, size=12, bold=False, cn="宋体", en="Times New Roman"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), en)
    fonts.set(qn("w:hAnsi"), en)
    fonts.set(qn("w:eastAsia"), cn)


def clear_cell_paragraphs(cell, keep=1):
    for extra in cell.paragraphs[keep:]:
        extra._p.getparent().remove(extra._p)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    return p


def normalize_fonts(doc):
    """模板中的 楷体_GB2312 在现代系统常缺失，替换为 楷体。"""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        for run in p.runs:
            rpr = run._element.rPr
            if rpr is None:
                continue
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                continue
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                if fonts.get(qn(attr)) == "楷体_GB2312":
                    fonts.set(qn(attr), "楷体")


def widen_task_column(table, grid):
    """微调列宽：任务分工列加宽，使其在保持一页版式的前提下排得下分工说明。"""
    cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for col, width in zip(cols, grid):
        col.set(qn("w:w"), str(width))

    for row in table.rows:
        position = 0
        for cell in row.cells:
            span_el = cell._tc.tcPr.find(qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            width = sum(grid[position:position + span])
            tc_w = cell._tc.tcPr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.tcPr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            position += span


def main():
    doc = Document(str(TEMPLATE))
    table = doc.tables[0]
    widen_task_column(table, [640, 1150, 520, 950, 1250, 2500, 1779])

    # 作品编号：替换红字说明为书写横线
    cell = table.rows[0].cells[3]
    p = clear_cell_paragraphs(cell)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run("作品编号："), size=12, cn="楷体")
    set_run(p.add_run("＿" * 14), size=12, cn="楷体")

    # 作品题目
    cell = table.rows[1].cells[3]
    p = clear_cell_paragraphs(cell)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run("作品题目："), size=11, cn="楷体")
    set_run(p.add_run(TITLE), size=11, bold=True, cn="楷体")

    # 任务分工（第 4~6 行，第 6 列）
    for index, task in enumerate(TASKS):
        row = table.rows[4 + index]
        cell = row.cells[5]
        p = clear_cell_paragraphs(cell)
        fmt = p.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.1
        fmt.first_line_indent = Pt(0)
        set_run(p.add_run(task), size=8.5, cn="宋体")

    normalize_fonts(doc)

    props = doc.core_properties
    props.title = "审批及签字页"
    props.subject = "全国密码技术竞赛作品设计报告"
    props.author = ""
    props.last_modified_by = ""

    doc.save(str(OUT))
    print("已生成：", OUT, OUT.stat().st_size, "B")


if __name__ == "__main__":
    main()
