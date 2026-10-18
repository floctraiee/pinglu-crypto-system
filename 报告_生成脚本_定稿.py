# -*- coding: utf-8 -*-
"""按官方模板生成《作品设计报告》定稿 DOCX。

用法：
    python 报告_生成脚本_定稿.py                          # 生成完整报告（目录页码取缓存值）
    python 报告_生成脚本_定稿.py --pages _tmp/toc_pages.json
    python 报告_生成脚本_定稿.py --measure-out _tmp/toc_measure   # 为测页码导出截断副本

设计要点：
    1. 直接以官方模板（附件二）为底稿，保留封面、基本信息表、页眉与页面设置；
    2. 正文只使用模板已有的 Heading 1 / Heading 2 样式，保证 Word 能重建目录；
    3. 目录写成 TOC 域，缓存条目带页码；并在 settings.xml 中置 updateFields，
       Word 打开时会自动重建目录（页码、超链接随实际分页刷新）。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import 报告_内容_定稿 as C  # noqa: E402

ROOT = Path(__file__).resolve().parent
# 官方模板不入库：用环境变量指定路径，避免把本机个人路径写进提交包
#   set PINGLU_REPORT_TEMPLATE=<附件二 模板.docx 路径>
TEMPLATE = Path(os.environ.get("PINGLU_REPORT_TEMPLATE")
                or (ROOT.parent / "附件二---作品设计报告（模板）.docx"))
FIG_DIR = ROOT / "报告图_定稿"
OUT = ROOT / "作品设计报告_面向平陆运河数字孪生的多源感知数据国密可信采集与篡改溯源系统.docx"

CN_HEI = "黑体"
CN_SONG = "宋体"
CN_KAI = "楷体_GB2312"
EN = "Times New Roman"

BODY_PT = 12.0        # 小四
TABLE_PT = 9.0        # 小五
CAPTION_PT = 10.5     # 五号
CONTENT_CM = 17.0     # A4 21cm - 左右各 2cm

INK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
HEAD_FILL = "DCE6F1"


# --------------------------------------------------------------------- 基础工具
def set_run(run, size=BODY_PT, bold=False, cn=CN_SONG, en=EN, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), en)
    fonts.set(qn("w:hAnsi"), en)
    fonts.set(qn("w:cs"), en)
    fonts.set(qn("w:eastAsia"), cn)
    return run


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def fld_char(kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    return el


def disable_doc_grid(doc):
    """模板带 w:docGrid（字符网格），部分渲染器会据此排版导致正文右移、右边距失效。"""
    for section in doc.sections:
        for grid in section._sectPr.findall(qn("w:docGrid")):
            grid.set(qn("w:type"), "none")
            for attr in ("w:charSpace",):
                if grid.get(qn(attr)) is not None:
                    del grid.attrib[qn(attr)]


def normalize_cover_fonts(doc):
    """模板封面用 楷体_GB2312（老字体，现代系统常缺失），统一替换为 楷体。"""
    def fix(paragraphs):
        for p in paragraphs:
            for run in p.runs:
                rpr = run._element.rPr
                if rpr is None:
                    continue
                fonts = rpr.find(qn("w:rFonts"))
                if fonts is None:
                    continue
                for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                    if fonts.get(qn(attr)) == "楷体_GB2312":
                        fonts.set(qn(attr), "楷体")

    fix(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                fix(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                fix(part.paragraphs)
            except Exception:
                pass


def set_update_fields(doc):
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def insert_pg_num_type(sect_pr, start=1):
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:start"), str(start))
    sect_pr.insert_element_before(
        el, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange",
    )
    return el


# --------------------------------------------------------------------- 正文构件
def body_para(doc, text, size=BODY_PT, indent=2.0, space_after=6, bold=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Pt(size * indent)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run(text), size=size, bold=bold)
    return p


def heading(doc, text, level):
    if level in (1, 2):
        p = doc.add_paragraph(style=f"Heading {level}")
    else:
        p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.line_spacing = 1.3
    fmt.keep_with_next = True
    fmt.keep_together = True
    if level == 1:
        fmt.space_before = Pt(16)
        fmt.space_after = Pt(9)
        set_run(p.add_run(text), size=15, bold=True, cn=CN_HEI)
    elif level == 2:
        fmt.space_before = Pt(11)
        fmt.space_after = Pt(6)
        set_run(p.add_run(text), size=13, bold=True, cn=CN_HEI)
    else:
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(4)
        set_run(p.add_run(text), size=12, bold=True, cn=CN_HEI)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(4)
    fmt.line_spacing = 1.15
    fmt.first_line_indent = Pt(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.keep_with_next = True
    set_run(p.add_run(text), size=CAPTION_PT, bold=True, cn=CN_SONG, color=INK)
    return p


def figure(doc, filename, width_cm, cap):
    path = FIG_DIR / filename
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(2)
    fmt.line_spacing = 1.0
    fmt.first_line_indent = Pt(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.keep_with_next = True
    if path.exists():
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        set_run(p.add_run(f"（缺图：{filename}）"), size=9, color=GREY)
    cap_p = caption(doc, cap)
    cap_p.paragraph_format.keep_with_next = False
    cap_p.paragraph_format.space_after = Pt(10)
    set_run_children_bold(cap_p, False)
    return p


def set_run_children_bold(p, bold):
    for r in p.runs:
        r.font.bold = bold


def tblpr_insert(tbl_pr, el, successors):
    tbl_pr.insert_element_before(el, *successors)


TBL_TAIL = (
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)


def apply_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tblpr_insert(tbl_pr, borders, ("w:shd", "w:tblLayout") + TBL_TAIL)


def add_table(doc, headers, rows, widths, cap):
    widths = list(widths)
    total = sum(widths)
    scale = CONTENT_CM / total
    widths = [w * scale for w in widths]

    caption(doc, cap)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr_insert(tbl_pr, layout, TBL_TAIL)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(Cm(CONTENT_CM).twips)))
    tblpr_insert(tbl_pr, tbl_w, ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                                 "w:shd", "w:tblLayout") + TBL_TAIL)
    apply_table_borders(table)

    for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(Cm(width).twips)))

    def fill_cell(cell, text, bold, align, size):
        cell.text = ""
        parts = str(text).split("\n")
        for index, part in enumerate(parts):
            p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            fmt = p.paragraph_format
            fmt.space_before = Pt(1.5)
            fmt.space_after = Pt(1.5)
            fmt.line_spacing = 1.1
            fmt.first_line_indent = Pt(0)
            fmt.alignment = align
            set_run(p.add_run(part), size=size, bold=bold, cn=CN_SONG)

    hdr_row = table.rows[0]
    tr_pr = hdr_row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    tr_pr.append(OxmlElement("w:cantSplit"))
    for index, title in enumerate(headers):
        cell = hdr_row.cells[index]
        cell.width = Cm(widths[index])
        fill_cell(cell, title, True, WD_ALIGN_PARAGRAPH.CENTER, TABLE_PT)
        shade(cell._tc.get_or_add_tcPr(), HEAD_FILL)

    for row in rows:
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(row):
            cells[index].width = Cm(widths[index])
            fill_cell(cells[index], value, False, WD_ALIGN_PARAGRAPH.LEFT, TABLE_PT)

    # 小表整体不跨页：除最后一行外，所有单元格段落与下一段同页
    if len(rows) <= 8:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(8)
    gap.paragraph_format.line_spacing = 1.0
    set_run(gap.add_run(""), size=6)
    return table


def bullet_list(doc, items, size=BODY_PT):
    for item in items:
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(4)
        fmt.line_spacing = 1.45
        fmt.left_indent = Cm(0.9)
        fmt.first_line_indent = Cm(-0.45)
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run(p.add_run("· "), size=size, bold=True)
        set_run(p.add_run(item), size=size)


def code_block(doc, lines, size=10.0):
    for item in lines:
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.25
        fmt.left_indent = Cm(0.9)
        fmt.first_line_indent = Cm(0)
        fmt.keep_together = True
        shade(p._p.get_or_add_pPr(), "F4F6F8")
        set_run(p.add_run(item), size=size, cn=CN_SONG, en="Consolas")
    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(8)
    tail.paragraph_format.line_spacing = 1.0
    set_run(tail.add_run(""), size=6)


# --------------------------------------------------------------------- 封面与基本信息表
def drop_text_runs(p):
    """删除段落里的文字 run，保留内嵌图形（如封面学会标识）。"""
    for r in list(p.runs):
        r_el = r._r
        if r_el.find(qn("w:drawing")) is not None or r_el.find(qn("w:pict")) is not None:
            continue
        r_el.getparent().remove(r_el)


def fill_cover(doc):
    paras = doc.paragraphs
    for p in paras:
        text = p.text
        if text.startswith("作品编号"):
            drop_text_runs(p)
            set_run(p.add_run("作品编号："), size=14, cn=CN_KAI)
            report_id = getattr(C, "REPORT_ID", "")
            set_run(p.add_run(report_id or ("＿" * 18)), size=14, cn=CN_KAI)
        elif text.startswith("作品类别"):
            drop_text_runs(p)
            set_run(p.add_run("作品类别： ■软件设计    □硬件制作    □工程实践"),
                    size=14, cn=CN_KAI)
        elif text.startswith("□密码应用技术"):
            drop_text_runs(p)
            set_run(p.add_run("■密码应用技术    □其它"), size=14, cn=CN_KAI)
        elif text.startswith("作品题目"):
            drop_text_runs(p)
            set_run(p.add_run("作品题目："), size=14, cn=CN_KAI)
            # 题目较长：在语义处断行，两行都对齐到题目起始位置，避免第二行回退到左边距
            fmt = p.paragraph_format
            fmt.left_indent = Pt(70)
            fmt.first_line_indent = Pt(-70)
            title_text = C.TITLE
            cut = 17 if len(title_text) > 24 else len(title_text)
            set_run(p.add_run(title_text[:cut]), size=14, bold=True, cn=CN_KAI)
            if cut < len(title_text):
                p.add_run().add_break(WD_BREAK.LINE)
                set_run(p.add_run(title_text[cut:]), size=14, bold=True, cn=CN_KAI)
        elif text.strip() == "年  月  日":
            drop_text_runs(p)
            set_run(p.add_run("2026年10月19日"), size=14, cn=CN_KAI)
        elif text.startswith("注意：作品为匿名评审"):
            p._p.getparent().remove(p._p)


def fill_info_table(doc):
    table = doc.tables[0]

    # 表格与页眉横线之间留一点空白
    spacer = doc.add_paragraph()
    table._tbl.addprevious(spacer._p)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    set_run(spacer.add_run(""), size=14)

    def rewrite(cell, segments):
        """segments: [(text, bold)]，清空单元格后按顺序写入。"""
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        fmt = p.paragraph_format
        fmt.space_before = Pt(2)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.2
        fmt.first_line_indent = Pt(0)
        for text, bold in segments:
            set_run(p.add_run(text), size=10.5, bold=bold, cn=CN_SONG)
        return p

    rewrite(table.rows[1].cells[0], [("编号：", False),
                                     (getattr(C, "REPORT_ID", "") or ("＿" * 16), False)])
    rewrite(table.rows[2].cells[0], [("作品题目：", False), (C.TITLE, True)])
    rewrite(table.rows[3].cells[0], [
        ("作品类别：", False),
        ("■软件设计    □硬件制作    □工程实践    ■密码技术应用    □其它", False),
    ])

    # 摘要：单元格原本由若干空段落构成书写行，这里替换为真实摘要
    cell = table.rows[4].cells[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    fmt = first.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(2)
    fmt.line_spacing = 1.25
    fmt.first_line_indent = Pt(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(first.add_run("作品内容摘要："), size=10.5, cn=CN_SONG)
    set_run(first.add_run(C.ABSTRACT), size=10.5, cn=CN_SONG)
    first.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    first.paragraph_format.first_line_indent = Pt(0)

    rewrite(table.rows[5].cells[0], [("关键词（五个）：", False), (C.KEYWORDS, False)])


# --------------------------------------------------------------------- 目录
def toc_entries():
    items = []
    for block in C.BLOCKS:
        if block[0] == "h1":
            items.append((1, block[1]))
        elif block[0] == "h2":
            items.append((2, block[1]))
    return items


def insert_toc(doc, anchor_p, pages):
    anchor_el = anchor_p._p if hasattr(anchor_p, "_p") else anchor_p

    def new_para():
        p = doc.add_paragraph()
        anchor_el.addprevious(p._p)
        return p

    title = new_para()
    fmt = title.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(16)
    fmt.line_spacing = 1.0
    fmt.page_break_before = True
    set_run(title.add_run("目　　录"), size=16, bold=True, cn=CN_HEI)

    entries = toc_entries()
    for index, (level, text) in enumerate(entries):
        p = new_para()
        fmt = p.paragraph_format
        fmt.space_before = Pt(4 if level == 1 else 0)
        fmt.space_after = Pt(4 if level == 1 else 2)
        fmt.line_spacing = 1.5 if level == 1 else 1.35
        fmt.first_line_indent = Pt(0)
        fmt.left_indent = Cm(0 if level == 1 else 1.0)
        fmt.tab_stops.add_tab_stop(Cm(CONTENT_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        if index == 0:
            r = p.add_run()
            r._r.append(fld_char("begin"))
            r2 = p.add_run()
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = ' TOC \\o "1-2" \\h \\z \\u '
            r2._r.append(instr)
            r3 = p.add_run()
            r3._r.append(fld_char("separate"))

        page = pages.get(text, 1)
        set_run(p.add_run(f"{text}\t{page}"), size=12, bold=(level == 1), cn=CN_SONG)

        if index == len(entries) - 1:
            r = p.add_run()
            r._r.append(fld_char("end"))


# --------------------------------------------------------------------- 正文
def build_body(doc, blocks):
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            heading(doc, block[1], 1)
        elif kind == "h2":
            heading(doc, block[1], 2)
        elif kind == "h3":
            heading(doc, block[1], 3)
        elif kind == "p":
            body_para(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1], block[2], block[3], block[4])
        elif kind == "figure":
            figure(doc, block[1], block[2], block[3])
        elif kind == "bullets":
            bullet_list(doc, block[1])
        elif kind == "code":
            code_block(doc, block[1])
        elif kind == "pagebreak":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(f"未知块类型：{kind}")


# --------------------------------------------------------------------- 组装
def delete_body(doc):
    """删除模板正文示例段落，保留第 2 页末尾的分节符与文档级 sectPr。"""
    body = doc.element.body
    children = list(body)
    anchor = None
    for el in children:
        if el.tag == qn("w:p") and el.find(qn("w:pPr")) is not None:
            if el.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
                anchor = el  # 第 2 节结束段落（最后一个带 sectPr 的 p）
    assert anchor is not None, "未找到模板分节符"
    started = False
    for el in children:
        if el is anchor:
            started = True
            continue
        if not started:
            continue
        if el.tag == qn("w:sectPr"):
            continue
        body.remove(el)
    return anchor


def drop_footer(section):
    """取消该节的页脚引用（模板页脚里带页码域，前置页不编号）。"""
    sect_pr = section._sectPr
    for ref in sect_pr.findall(qn("w:footerReference")):
        sect_pr.remove(ref)


def add_footer_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("— "), size=9, color=GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), EN)
    fonts.set(qn("w:hAnsi"), EN)
    fonts.set(qn("w:eastAsia"), CN_SONG)
    rpr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "595959")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)
    set_run(p.add_run(" —"), size=9, color=GREY)


def build(out_path, pages=None, limit=None):
    pages = pages or {}
    doc = Document(str(TEMPLATE))
    fill_cover(doc)
    fill_info_table(doc)

    anchor = delete_body(doc)
    insert_toc(doc, anchor, pages)

    blocks = C.BLOCKS if limit is None else C.BLOCKS[:limit]
    build_body(doc, blocks)
    if limit is not None:
        # 供测页码使用：补一行正文，避免标题带 keep_with_next 时被判成不需要换页
        body_para(doc, "　")

    body_section = doc.sections[-1]
    insert_pg_num_type(body_section._sectPr, start=1)
    add_footer_page_number(body_section)

    disable_doc_grid(doc)
    normalize_cover_fonts(doc)
    # 必须在 normalize 之后取消前置页页脚：访问 section.footer 可能补建页脚定义
    for section in doc.sections[:-1]:
        drop_footer(section)
    set_update_fields(doc)

    props = doc.core_properties
    props.title = C.TITLE
    props.subject = "全国密码技术竞赛作品设计报告"
    props.keywords = C.KEYWORDS
    props.author = "匿名"
    props.last_modified_by = "匿名"
    props.comments = ""
    props.category = "软件设计"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(OUT))
    parser.add_argument("--pages", default=None, help="目录页码 JSON：{标题: 页码}")
    parser.add_argument("--measure-out", default=None, help="导出测页码用的截断副本目录")
    args = parser.parse_args()

    if args.measure_out:
        target = Path(args.measure_out)
        target.mkdir(parents=True, exist_ok=True)
        for index, block in enumerate(C.BLOCKS):
            if block[0] in ("h1", "h2"):
                path = target / f"m{index:03d}.docx"
                build(path, pages={}, limit=index + 1)
        build(target / "m000_front.docx", pages={}, limit=0)
        print("已导出测页码副本：", len(list(target.glob('*.docx'))), "个")
        return

    pages = {}
    if args.pages:
        pages = json.loads(Path(args.pages).read_text(encoding="utf-8"))
    path = build(Path(args.out), pages=pages)
    print("已生成：", path, path.stat().st_size, "B")
    print("目录条目：", len(toc_entries()), "条；页码来自缓存：", bool(pages))


if __name__ == "__main__":
    main()
